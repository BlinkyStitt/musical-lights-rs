"""Compare the Rust stream with unchanged ISO fixtures and an independent oracle."""

import argparse
import hashlib
import json
import os
import shutil
import subprocess
import urllib.request
import warnings
import zipfile
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.io import wavfile

ISO_URL = (
    "https://standards.iso.org/iso/532/-1/ed-1/en/ISO%20532-1%20-%20Program%20etc.zip"
)
ISO_SHA256 = "d17b2c6d66a28550ed145c3e1ae5af6ee5917b90e584358285686b1fc61edca7"
ROOT = Path(__file__).resolve().parents[2]
FRAME_DTYPE = np.dtype(
    [
        ("sample", "<u8"),
        ("sones", "<f8"),
        ("specific", "<f8", (240,)),
    ]
)


def limits(actual, reference):
    error = np.abs(actual - reference)
    outer = error > np.maximum(0.2, reference * 0.1)
    inner = error > np.maximum(0.1, reference * 0.05)
    return {
        "max_absolute_error": float(error.max()),
        "outside_outer": int(outer.sum()),
        "outside_inner_fraction": float(inner.mean()),
        "passed": bool(not outer.any() and inner.mean() <= 0.01),
    }


def run_trace(binary, pressure, cache, label, field="free", block=128):
    source = cache / f"{label}.f32"
    target = cache / f"{label}.bin"
    np.asarray(pressure, dtype="<f4").tofile(source)
    subprocess.run(
        [str(binary), str(source), str(target), str(block), "1", field],
        check=True,
        timeout=120,
    )
    result = np.fromfile(target, dtype=FRAME_DTYPE)
    expected = np.arange(0, len(pressure), 96, dtype=np.uint64)
    np.testing.assert_array_equal(result["sample"], expected)
    assert np.isfinite(result["sones"]).all() and np.isfinite(result["specific"]).all()
    assert (result["sones"] >= 0).all() and (result["specific"] >= 0).all()
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cache", type=Path, default=ROOT / ".cache/loudness")
    parser.add_argument("--archive", type=Path)
    parser.add_argument("--cases", type=int, nargs="+", default=list(range(6, 26)))
    parser.add_argument("--oracle", action="store_true")
    parser.add_argument(
        "--binary", type=Path, default=ROOT / "target/release/examples/loudness_trace"
    )
    args = parser.parse_args()
    cache = args.cache.resolve()
    cache.mkdir(parents=True, exist_ok=True)
    archive = args.archive or cache / "iso-532-1.zip"
    if not archive.exists():
        urllib.request.urlretrieve(ISO_URL, archive)
    digest = hashlib.sha256(archive.read_bytes()).hexdigest()
    if digest != ISO_SHA256:
        raise ValueError(f"ISO archive checksum mismatch: {digest}")
    fixtures = cache / "iso"
    # Rebuild only this validator's extracted fixtures from the verified archive.
    # A previous cache edit must not silently change the reference measurements.
    if fixtures.exists():
        shutil.rmtree(fixtures)
    with zipfile.ZipFile(archive) as contents:
        contents.extractall(fixtures)
    os.environ.setdefault("MPLCONFIGDIR", str(cache / "matplotlib"))
    import matplotlib

    matplotlib.use("Agg")
    from matplotlib import pyplot as plt

    report = {"iso_archive_sha256": digest, "cases": [], "hardware": "pending"}
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore", message="Conditional Formatting extension is not supported.*"
        )
        books = {
            annex: load_workbook(
                next((fixtures / f"Annex B.{annex}").glob("*.xlsx")),
                read_only=True,
                data_only=True,
            )
            for annex in (4, 5)
        }
    for case in args.cases:
        annex = 4 if case < 14 else 5
        wav = next((fixtures / f"Annex B.{annex}").glob(f"Test signal {case} (*.wav"))
        rate, pcm = wavfile.read(wav)
        assert rate == 48000 and pcm.dtype == np.int16 and pcm.ndim == 1
        pressure = (pcm.astype(np.float64) * (2 * np.sqrt(2) / 32768)).astype(
            np.float32
        )
        field = "diffuse" if case == 15 else "free"
        actual = run_trace(args.binary, pressure, cache, f"iso-{case}", field)
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore", message="Conditional Formatting extension is not supported.*"
            )
            rows = [
                r
                for r in books[annex][f"Test signal {case}"].iter_rows(
                    min_row=11, values_only=True
                )
                if isinstance(r[0], (int, float))
            ]
        time = np.asarray([r[0] for r in rows])
        # Some worksheets include the endpoint at duration, where no input sample
        # exists. Match by published timestamps, never by best-fit time shifting.
        keep = time < len(pressure) / rate
        rows = [r for r, include in zip(rows, keep, strict=True) if include]
        time = time[keep]
        # The technical-sound worksheets omit a final incomplete 2 ms interval.
        # Check every published timestamp and keep the extra emitted frame in
        # the stream/oracle checks. No calculated frame is moved in time.
        assert len(time) in {len(pressure) // 96, (len(pressure) + 95) // 96}
        compared = actual[: len(time)]
        np.testing.assert_allclose(compared["sample"] / rate, time, rtol=0, atol=1e-9)
        reference = np.asarray([r[1] for r in rows])
        result = {
            "case": case,
            "field": field,
            "unreferenced_final_frames": len(actual) - len(compared),
            "total": limits(compared["sones"], reference),
        }
        if case < 14:
            bark = {6: 2.5, 8: 17.5, 9: 17.5}.get(case, 8.5)
            specific = np.asarray([r[11] for r in rows])
            result["specific"] = limits(
                compared["specific"][:, round(bark * 10) - 1], specific
            )
        if args.oracle:
            from mosqito.sq_metrics import loudness_zwtv

            oracle, oracle_specific, _, _ = loudness_zwtv(
                pressure.astype(np.float64), rate, field_type=field
            )
            result["oracle_total_max_error"] = float(
                np.max(np.abs(actual["sones"] - oracle))
            )
            result["oracle_specific_max_error"] = float(
                np.max(np.abs(actual["specific"] - oracle_specific.T))
            )
            result["oracle_total"] = limits(actual["sones"], oracle)
            result["oracle_specific"] = limits(
                actual["specific"].reshape(-1), oracle_specific.T.reshape(-1)
            )
        report["cases"].append(result)
        print(json.dumps(result), flush=True)
        fig, ax = plt.subplots(figsize=(9, 3))
        ax.plot(time, reference, label="ISO reference", color="#4477AA")
        ax.plot(time, compared["sones"], label="Rust", color="#EE6677", linestyle="--")
        ax.set(xlabel="Audio time (s)", ylabel="Loudness (sone)", title=wav.stem)
        ax.legend()
        fig.tight_layout()
        fig.savefig(cache / f"iso-{case}.svg")
        plt.close(fig)
    for book in books.values():
        book.close()
    (cache / "report.json").write_text(json.dumps(report, indent=2) + "\n")
    failed = [
        r["case"]
        for r in report["cases"]
        if any(
            not r.get(key, {"passed": True})["passed"]
            for key in ("total", "specific", "oracle_total", "oracle_specific")
        )
    ]
    if failed:
        raise SystemExit(f"ISO reference failures: {failed}")


if __name__ == "__main__":
    main()
